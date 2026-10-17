#!/usr/bin/env python3
"""Simple teacher list with address — API-based, fast."""
import asyncio
import json
import os
import re
from datetime import datetime
from pathlib import Path
from playwright.async_api import async_playwright
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side

BASE = "https://eshikshakosh.bihar.gov.in"
API = "https://eshikshakosh.bihar.gov.in:8443"
AUTH_HEADER_NAMES = {"accept", "content-type", "authorization", "accesstoken", "authtoken", "authaccess"}
GENDER_MAP = {"1": "Male", "2": "Female"}


def require_env(name):
    val = (os.environ.get(name) or "").strip()
    if not val:
        raise RuntimeError(f"{name} is not set")
    return val


def unpack(response_json):
    if not isinstance(response_json, dict):
        return [], 0
    rows = response_json.get("data", [])
    total = response_json.get("totalRecord", 0)
    if isinstance(rows, dict):
        total = rows.get("totalRecord", total)
        rows = rows.get("data", [])
    if not isinstance(rows, list):
        rows = []
    try:
        total = int(total or len(rows))
    except (ValueError, TypeError):
        total = len(rows)
    return rows, total


async def main():
    uid = require_env("ESHIKSHAKOSH_USERNAME")
    pwd = require_env("ESHIKSHAKOSH_PASSWORD")
    target_udise = require_env("ESHIKSHAKOSH_TARGET_UDISE")
    out_dir = Path(os.environ.get("ESHIKSHAKOSH_OUT") or (Path.home() / "exports"))

    print(f"School UDISE: {target_udise}")

    auth = {}
    teachers, total = [], 0

    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True, args=["--no-sandbox", "--disable-dev-shm-usage"])
        context = await browser.new_context(ignore_https_errors=True)
        page = await context.new_page()

        async def capture_auth(route):
            req = route.request
            try:
                payload = json.loads(req.post_data or "{}")
            except:
                await route.continue_()
                return
            payload["offset"] = "0"
            payload["limit"] = "100"
            payload["schoolUdiseCode"] = target_udise
            hdrs = await req.all_headers()
            auth.update({k.lower(): v for k, v in hdrs.items() if k.lower() in AUTH_HEADER_NAMES})
            await route.continue_(post_data=json.dumps(payload))

        await page.route("**/teacher/viewTeacher", capture_auth)

        try:
            # Login
            await page.goto(BASE, wait_until="networkidle", timeout=90000)
            await page.wait_for_timeout(3000)
            if await page.locator("#radio2").count():
                await page.click("#radio2")
                await page.wait_for_timeout(300)
            await page.locator("input[formcontrolname='userId']").wait_for(state="visible", timeout=15000)
            await page.fill("input[formcontrolname='userId']", uid)
            await page.fill("input[formcontrolname='password']", pwd)

            captcha = await page.inner_text(".capcha__head")
            m = re.search(r"(\d+)\s*([+\-*])\s*(\d+)", captcha)
            if not m:
                raise RuntimeError("Captcha not recognized")
            a, op, b = int(m.group(1)), m.group(2), int(m.group(3))
            await page.fill("input[placeholder='Enter User Captcha']", str(a + b if op == "+" else a - b if op == "-" else a * b))

            async with page.expect_response(lambda r: "/auth/login" in r.url.lower(), timeout=35000) as login_info:
                await page.click("input.submit__btn")
            login_resp = await login_info.value
            print(f"Login: HTTP {login_resp.status}")
            if login_resp.status != 200:
                raise RuntimeError(f"Login failed")
            await page.wait_for_timeout(1200)

            # Get teacher list
            await page.goto(f"{BASE}/Application/teacher/registration/viewTeacher", wait_until="networkidle", timeout=60000)
            await page.wait_for_timeout(2000)
            search = page.locator("button:has-text('Search')")
            if not await search.count():
                raise RuntimeError("Search button not found")
            async with page.expect_response(lambda r: "/teacher/viewTeacher" in r.url and r.request.method == "POST", timeout=45000) as teacher_info:
                await search.first.click()
            first_json = await (await teacher_info.value).json()
            teachers, total = unpack(first_json)
            if not teachers:
                raise RuntimeError("No teachers found")
            if not auth:
                raise RuntimeError("Could not capture auth")

            # Pagination
            while len(teachers) < total:
                offset = len(teachers)
                payload = {"offset": str(offset), "limit": "100", "schoolUdiseCode": target_udise}
                res = await page.evaluate("""async ({url, payload, headers}) => {
                    const r = await fetch(url, {method:'POST', headers, body:JSON.stringify(payload)});
                    let body; try { body = await r.json(); } catch (_) { body = null; }
                    return {status:r.status, body};
                }""", {"url": f"{API}/teacher/viewTeacher", "payload": payload, "headers": auth})
                if res.get("status") != 200:
                    raise RuntimeError(f"Pagination failed")
                nxt, nxt_total = unpack(res.get("body"))
                if not nxt:
                    break
                teachers.extend(nxt)
                total = max(total, nxt_total)

            print(f"Teachers mile: {len(teachers)}")

            # Fetch address for each teacher
            print("Address fetch ho raha hai...")
            addresses = {}
            for idx, teacher in enumerate(teachers, 1):
                enc = str(teacher.get("encId") or "").strip()
                name = teacher.get("teacherName", "")
                if enc:
                    try:
                        res = await page.evaluate("""async ({url, payload, headers}) => {
                            const r = await fetch(url, {method:'POST', headers, body:JSON.stringify(payload)});
                            let body; try { body = await r.json(); } catch (_) { body = null; }
                            return {status:r.status, body};
                        }""", {"url": f"{API}/teacher/registrationInfo", "payload": {"encId": enc}, "headers": auth})
                        if res.get("status") == 200 and isinstance(res.get("body"), dict):
                            addresses[enc] = res["body"].get("data", {})
                    except:
                        pass
                print(f"  [{idx}/{len(teachers)}] {name}")
                await page.wait_for_timeout(300)

        finally:
            await context.close()
            await browser.close()

    # Export to Excel
    out_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    school_name = teachers[0].get("schoolName", "Unknown") if teachers else "Unknown"
    school_slug = re.sub(r"[^A-Za-z0-9]+", "_", school_name).strip("_")
    output = out_dir / f"Teacher_List_{school_slug}_{timestamp}.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Teachers"

    headers = ["S.No.", "Temp Login ID", "Teacher Name", "Gender", "Marital Status",
               "DOB", "Teacher Title", "Nature of Appointment", "Service Joining",
               "Joining Current School", "Class Category", "Appointed Subject",
               "Mobile", "Email", "PAN", "HRMS ID", "PRAN", "GPF", "UAN",
               "Aadhaar", "Social Category", "Domicile No", "Mother", "Father",
               "Current Address", "Permanent Address", "School Name", "UDISE"]
    ws.append(headers)

    # Formatting
    header_fill = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid")
    header_font = Font(bold=True, color="1F2937", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )
    for c in ws[1]:
        c.fill = header_fill
        c.font = header_font
        c.border = thin_border

    for i, t in enumerate(teachers, 1):
        enc = str(t.get("encId") or "").strip()
        addr = addresses.get(enc, {})

        # Build combined address
        curr_parts = []
        for k in ["cHouseName", "cStreet", "cPanchayat", "cVillage", "cPIN", "cDistrictId", "cBlockId", "cStateName"]:
            v = addr.get(k, "")
            if v and str(v).strip():
                curr_parts.append(str(v).strip())
        curr_addr = ", ".join(curr_parts) if curr_parts else ""

        perm_parts = []
        for k in ["pHouseName", "pStreet", "pPanchayat", "pVillage", "pPIN", "pDistrictId", "pBlockId", "pStateName"]:
            v = addr.get(k, "")
            if v and str(v).strip():
                perm_parts.append(str(v).strip())
        perm_addr = ", ".join(perm_parts) if perm_parts else ""

        def label(key, fallback=""):
            """Readable label from registrationInfo; falls back to raw list value."""
            v = addr.get(key)
            if v in (None, "", "--"):
                return fallback
            return v

        # gender: prefer readable label from registrationInfo
        gender = label("gender") or t.get("gender", "")
        if str(gender).strip() in GENDER_MAP:
            gender = GENDER_MAP[str(gender).strip()]

        ws.append([
            i, t.get("teacherLoginCode"), t.get("teacherName"), gender,
            label("maritalStatus", t.get("maritalStatus")), t.get("DOB"), t.get("teacherTitle"),
            label("natureOfAppointmt", t.get("natureOfAppointmt")), t.get("serviceJoiningDt"),
            t.get("joiningCurrentSchoolDt"), label("classCategory", t.get("classCategory")),
            label("appointedSubject", t.get("appointedSubject")),
            # mobile / PAN / Aadhaar stay unmasked -> must come from the list API
            t.get("mobile"), t.get("teacherEmail"),
            t.get("panNo"), t.get("employeeCode"), t.get("pranNo"),
            t.get("gpfNo") or "", t.get("uanNo") or "", t.get("aadhaarNo"),
            label("socialCategory", t.get("socialCategory")), t.get("pDomicileNumber"),
            t.get("motherName"), t.get("fatherName"),
            curr_addr, perm_addr,
            t.get("schoolName"), t.get("schoolUdiseCode"),
        ])

    # Auto-size columns
    from openpyxl.utils import get_column_letter
    for col in range(1, len(headers) + 1):
        max_len = len(str(headers[col-1]))
        for row in range(2, min(ws.max_row, 60) + 1):
            val = ws.cell(row=row, column=col).value
            if val:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[get_column_letter(col)].width = min(max_len + 2, 40)

    ws.freeze_panes = "A2"
    wb.save(output)
    os.chmod(output, 0o600)

    print(f"\nTeachers: {len(teachers)}")
    print(f"File: {output}")


if __name__ == "__main__":
    asyncio.run(main())
