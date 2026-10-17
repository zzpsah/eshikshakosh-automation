#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
EshikshaKosh Teacher Report — Playwright-based extraction.

Logs in via Playwright, navigates to teacher registration page,
clicks each teacher's "Print Teacher Details" badge (opens new tab),
extracts all content from the print view, and exports to Excel.
"""

import asyncio
import logging
import re
import sys
import time
from pathlib import Path

from playwright.async_api import async_playwright
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
log = logging.getLogger("esk_teacher_pw")


def safe_filename_component(value: str, fallback: str = "School") -> str:
    cleaned = "".join(ch if ch.isalnum() or ch in (" ", "-", "_") else " " for ch in str(value or ""))
    cleaned = "_".join(cleaned.split()).strip("._-")
    return cleaned or fallback


async def login_and_extract(uid: str, pwd: str, output: str, verify_ssl=False):
    """Login and extract teacher details via Print Teacher Details popups."""

    async with async_playwright() as p:
        browser = await p.chromium.launch(
            headless=True,
            args=["--no-sandbox", "--disable-setuid-sandbox", "--disable-dev-shm-usage",
                  "--disable-popup-blocking"]
        )
        context = await browser.new_context(ignore_https_errors=True)
        page = await context.new_page()

        # === Login ===
        log.info("Navigating to login page...")
        await page.goto("https://eshikshakosh.bihar.gov.in",
                        wait_until="networkidle", timeout=35000)
        await page.click("#radio2")
        await page.wait_for_timeout(400)
        await page.fill("input[formcontrolname='userId']", uid)
        await page.fill("input[formcontrolname='password']", pwd)

        # Captcha
        captcha_text = await page.inner_text(".capcha__head")
        match = re.search(r"(\d+)\s*([+\-*])\s*(\d+)", captcha_text)
        if not match:
            raise RuntimeError("Captcha expression not recognised")
        left, op, right = int(match.group(1)), match.group(2), int(match.group(3))
        answer = left + right if op == "+" else left - right if op == "-" else left * right
        log.info("Auto-solved captcha: %d %s %d = %d", left, op, right, answer)
        await page.fill("input[placeholder='Enter User Captcha']", str(answer))

        async with page.expect_response(
            lambda resp: "/auth/login" in resp.url.lower(),
            timeout=25000,
        ) as login_info:
            await page.click("input.submit__btn")
        login_resp = await login_info.value
        if login_resp.status != 200:
            raise RuntimeError(f"Login failed with HTTP {login_resp.status}")

        # Wait for dashboard
        await page.wait_for_url("**/dashboard", timeout=10000)
        log.info("Login successful — on dashboard")

        # === Navigate to teacher registration page ===
        log.info("Navigating to teacher registration page...")
        await page.goto("https://eshikshakosh.bihar.gov.in/Application/teacher/registration/viewTeacher",
                        wait_until="networkidle", timeout=35000)
        await page.wait_for_timeout(2000)

        # Click Search to load teacher list
        search_btn = page.locator("button:has-text('Search'), button:has-text('खोज')").first
        await search_btn.click()
        await page.wait_for_timeout(3000)

        # === Extract teacher list from table ===
        all_teachers = []

        for page_num in range(2):  # Page 1 and Page 2
            if page_num > 0:
                # Click next page
                next_btn = page.locator(".mat-paginator-navigation-next, [aria-label='Next page']").first
                await next_btn.click()
                await page.wait_for_timeout(3000)

            rows = page.locator("table tbody tr")
            count = await rows.count()
            log.info("Page %d: %d teacher rows", page_num + 1, count)

            for i in range(count):
                row = rows.nth(i)
                cells = row.locator("td")
                cell_count = await cells.count()

                # Extract basic info from table
                sl = await cells.nth(1).text_content() if cell_count > 1 else ""
                temp_id = await cells.nth(2).text_content() if cell_count > 2 else ""
                teacher_id = await cells.nth(3).text_content() if cell_count > 3 else ""
                name = await cells.nth(4).text_content() if cell_count > 4 else ""
                teacher_type = await cells.nth(5).text_content() if cell_count > 5 else ""
                mobile = await cells.nth(6).text_content() if cell_count > 6 else ""

                teacher = {
                    "sl": sl.strip() if sl else "",
                    "temp_login_id": temp_id.strip() if temp_id else "",
                    "teacher_id": teacher_id.strip() if teacher_id else "",
                    "name": name.strip() if name else "",
                    "type": teacher_type.strip() if teacher_type else "",
                    "mobile": mobile.strip() if mobile else "",
                }
                all_teachers.append(teacher)
                log.info("  Found: %s", teacher["name"])

        log.info("Total teachers found: %d", len(all_teachers))

        # === Now click each teacher's Print Teacher Details and extract from popup ===
        log.info("Extracting full details from Print Teacher Details popups...")

        for idx, teacher in enumerate(all_teachers):
            # Navigate back to page 1 for first 10, page 2 for rest
            target_page = 1 if idx < 10 else 2
            current_page_num = 1 if idx == 0 else (2 if idx == 10 else current_page_num)

            if idx == 0:
                # Already on page 1
                pass
            elif idx == 10:
                # Navigate to page 2
                log.info("Navigating to page 2...")
                await page.goto("https://eshikshakosh.bihar.gov.in/Application/teacher/registration/viewTeacher",
                                wait_until="networkidle", timeout=35000)
                await page.wait_for_timeout(2000)
                search_btn = page.locator("button:has-text('Search'), button:has-text('खोज')").first
                await search_btn.click()
                await page.wait_for_timeout(3000)
                next_btn = page.locator(".mat-paginator-navigation-next, [aria-label='Next page']").first
                await next_btn.click()
                await page.wait_for_timeout(3000)
            elif idx > 0 and idx < 10:
                # Need to reload page 1 for each teacher (popup navigated away)
                if idx == 1:
                    log.info("Reloading page 1 for teacher popups...")
                    await page.goto("https://eshikshakosh.bihar.gov.in/Application/teacher/registration/viewTeacher",
                                    wait_until="networkidle", timeout=35000)
                    await page.wait_for_timeout(2000)
                    search_btn = page.locator("button:has-text('Search'), button:has-text('खोज')").first
                    await search_btn.click()
                    await page.wait_for_timeout(3000)

            # Find the row with this teacher's name and click its Print badge
            teacher_name = teacher["name"]
            log.info("[%d/%d] Clicking Print Details for: %s", idx + 1, len(all_teachers), teacher_name)

            # Find the row containing this teacher's name
            row_locator = page.locator(f"table tbody tr:has-text('{teacher_name}')").first

            # Click the Print Teacher Details badge in that row
            print_badge = row_locator.locator("a.badge:has-text('Print Teacher'), a.badge:has-text('शिक्षक विवरण')").first

            try:
                async with page.expect_popup(timeout=10000) as popup_info:
                    await print_badge.click()
                popup = await popup_info.value
                await popup.wait_for_load_state("networkidle", timeout=15000)
                await popup.wait_for_timeout(2000)

                # Extract content from popup
                popup_text = await popup.evaluate("document.body.innerText")

                # Parse the teacher details from popup text
                details = parse_teacher_details(popup_text)
                teacher["details"] = details

                log.info("    Extracted %d fields from popup", len(details))

                # Close popup
                await popup.close()
                await page.wait_for_timeout(1000)

                # Reload the page for next teacher (page state may have changed)
                if idx < len(all_teachers) - 1:
                    next_page = 1 if (idx + 1) < 10 else 2
                    current_target = 1 if idx < 10 else 2
                    if next_page != current_target:
                        pass  # Will handle navigation at the top of loop
                    else:
                        await page.goto("https://eshikshakosh.bihar.gov.in/Application/teacher/registration/viewTeacher",
                                        wait_until="networkidle", timeout=35000)
                        await page.wait_for_timeout(2000)
                        search_btn = page.locator("button:has-text('Search'), button:has-text('खोज')").first
                        await search_btn.click()
                        await page.wait_for_timeout(3000)
                        if current_target == 2:
                            next_btn = page.locator(".mat-paginator-navigation-next, [aria-label='Next page']").first
                            await next_btn.click()
                            await page.wait_for_timeout(3000)

            except Exception as e:
                log.warning("    Failed to extract popup for %s: %s", teacher_name, e)
                teacher["details"] = {}

        # === Build Excel ===
        log.info("Building Excel file...")
        build_excel(all_teachers, output)

        await page.close()
        await context.close()
        await browser.close()


def parse_teacher_details(text: str) -> dict:
    """Parse teacher details from popup text content."""
    details = {}
    lines = [l.strip() for l in text.split("\n") if l.strip()]

    # Common field patterns in the print view
    field_patterns = {
        "Teacher Name": r"Teacher Name\s*[:\-]?\s*(.+)",
        "Father's Name": r"Father'?s Name\s*[:\-]?\s*(.+)",
        "Mother's Name": r"Mother'?s Name\s*[:\-]?\s*(.+)",
        "Gender": r"Gender\s*[:\-]?\s*(.+)",
        "Date of Birth": r"(?:DOB|Date of Birth)\s*[:\-]?\s*(.+)",
        "Aadhaar Number": r"(?:Aadhaar|Aadhar)\s*(?:Number|No)?\s*[:\-]?\s*(\d+)",
        "Mobile Number": r"(?:Mobile|Phone)\s*(?:Number|No)?\s*[:\-]?\s*(\d+)",
        "Email": r"(?:Email|E-mail)\s*(?:Id|ID)?\s*[:\-]?\s*(.+)",
        "Address": r"Address\s*[:\-]?\s*(.+)",
        "District": r"District\s*[:\-]?\s*(.+)",
        "Block": r"Block\s*[:\-]?\s*(.+)",
        "Panchayat": r"Panchayat\s*[:\-]?\s*(.+)",
        "PIN Code": r"(?:PIN|Pincode)\s*[:\-]?\s*(.+)",
        "Category": r"Category\s*[:\-]?\s*(.+)",
        "Religion": r"Religion\s*[:\-]?\s*(.+)",
        "Caste": r"Caste\s*[:\-]?\s*(.+)",
        "Sub Caste": r"Sub\s*Caste\s*[:\-]?\s*(.+)",
        "Qualification": r"Qualification\s*[:\-]?\s*(.+)",
        "Subject": r"Subject\s*[:\-]?\s*(.+)",
        "Joining Date": r"(?:Joining|Date of Joining)\s*[:\-]?\s*(.+)",
        "Confirmation Date": r"Confirmation\s*Date\s*[:\-]?\s*(.+)",
        "Retirement Date": r"Retirement\s*Date\s*[:\-]?\s*(.+)",
        "Bank Account No": r"(?:Bank\s*)?Account\s*(?:No|Number)?\s*[:\-]?\s*(.+)",
        "IFSC Code": r"IFSC\s*[:\-]?\s*(.+)",
        "Bank Name": r"Bank\s*Name\s*[:\-]?\s*(.+)",
        "Branch": r"Branch\s*(?:Name)?\s*[:\-]?\s*(.+)",
        "GPF Number": r"GPF\s*(?:Number|No)?\s*[:\-]?\s*(.+)",
        "PRAN Number": r"PRAN\s*(?:Number|No)?\s*[:\-]?\s*(.+)",
        "UAN Number": r"UAN\s*(?:Number|No)?\s*[:\-]?\s*(.+)",
        "Nature of Appointment": r"Nature\s*of\s*Appointment\s*[:\-]?\s*(.+)",
        "Type of Teacher": r"Type\s*of\s*Teacher\s*[:\-]?\s*(.+)",
        "Temporary Login ID": r"Temporary\s*Login\s*Id\s*[:\-]?\s*(.+)",
        "Teacher ID": r"Teacher\s*ID\s*[:\-]?\s*(.+)",
        "Reference Number": r"(?:Reference\s*Number|Letter\s*Number|विषय)\s*[:\-]?\s*(.+)",
        "Issue Date": r"Issue\s*Date\s*[:\-]?\s*(.+)",
        "Authority": r"Authority\s*[:\-]?\s*(.+)",
    }

    for field, pattern in field_patterns.items():
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            details[field] = match.group(1).strip()

    return details


def build_excel(teachers: list, output: str):
    """Build Excel file from teacher data."""

    # Collect all unique detail keys
    all_detail_keys = set()
    for t in teachers:
        all_detail_keys.update(t.get("details", {}).keys())

    columns = [
        "S.No.", "Teacher Name", "Teacher ID", "Temp Login ID",
        "Type", "Mobile",
    ] + sorted(all_detail_keys)

    wb = Workbook()
    ws = wb.active
    ws.title = "Teachers"

    # Header
    hfill = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid")
    hfont = Font(name="Calibri", size=10, bold=True, color="1F2937")
    cfont = Font(name="Calibri", size=10, color="111827")
    border = Border(left=Side(style="thin", color="D1D5DB"),
                    right=Side(style="thin", color="D1D5DB"),
                    top=Side(style="thin", color="D1D5DB"),
                    bottom=Side(style="thin", color="D1D5DB"))

    ws.append(columns)
    ws.row_dimensions[1].height = 25
    for cell in ws[1]:
        cell.fill = hfill
        cell.font = hfont
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    # Data rows
    TEXT_COLS = {"Teacher ID", "Temp Login ID", "Aadhaar Number", "Mobile Number",
                 "Bank Account No", "IFSC Code", "GPF Number", "PRAN Number",
                 "UAN Number", "Reference Number"}
    CENTER_COLS = {"S.No.", "Gender", "Date of Birth", "Joining Date",
                   "Confirmation Date", "Retirement Date", "Issue Date"}

    text_idx = [columns.index(c) + 1 for c in TEXT_COLS if c in columns]
    center_idx = [columns.index(c) + 1 for c in CENTER_COLS if c in columns]

    for ri, teacher in enumerate(teachers, 2):
        ws.row_dimensions[ri].height = 20
        row_data = {
            "S.No.": ri - 1,
            "Teacher Name": teacher.get("name", ""),
            "Teacher ID": teacher.get("teacher_id", ""),
            "Temp Login ID": teacher.get("temp_login_id", ""),
            "Type": teacher.get("type", ""),
            "Mobile": teacher.get("mobile", ""),
        }
        # Add detail fields
        details = teacher.get("details", {})
        for key in all_detail_keys:
            row_data[key] = details.get(key, "")

        ws.append([row_data.get(c, "") for c in columns])
        for ci in range(1, len(columns) + 1):
            cell = ws.cell(row=ri, column=ci)
            cell.font = cfont
            cell.border = border
            if ci in text_idx:
                cell.number_format = "@"
                if cell.value is not None:
                    cell.value = str(cell.value).strip()
            cell.alignment = Alignment(horizontal="center", vertical="center") if ci in center_idx else Alignment(horizontal="left", vertical="center")

    # Column widths
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[letter].width = min(max(max_len + 4, 13), 45)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(output)
    log.info("Saved: %s", output)
    print(f"\n{'='*60}")
    print(f"  Report: {output}")
    print(f"  Total Teachers: {len(teachers)}")
    details_count = sum(1 for t in teachers if t.get("details"))
    print(f"  With Full Details: {details_count}")
    print(f"  Basic Info Only: {len(teachers) - details_count}")
    print(f"{'='*60}")


async def main():
    import argparse
    import os

    parser = argparse.ArgumentParser(description="EshikshaKosh Teacher Report (Playwright-based)")
    parser.add_argument("--udise", "-u", default=None)
    parser.add_argument("--password", "-p", default=None)
    parser.add_argument("--output", "-o", default=None)
    args = parser.parse_args()

    if not args.udise:
        args.udise = os.environ.get("ESHIKSHAKOSH_USERNAME", "")
    if not args.password:
        args.password = os.environ.get("ESHIKSHAKOSH_PASSWORD", "")

    if not args.udise:
        parser.error("--udise is required")
    if not args.password:
        import getpass
        args.password = getpass.getpass("Password: ")

    output = args.output or f"Teacher_Details_UMV_Tetahali.xlsx"

    await login_and_extract(args.udise, args.password, output)


if __name__ == "__main__":
    asyncio.run(main())
